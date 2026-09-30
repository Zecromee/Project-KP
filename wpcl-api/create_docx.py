from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(10, 22, 40)
    return h

def add_para(text, bold=False, italic=False, size=None, color=None, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    return p

def add_step(num, text):
    p = doc.add_paragraph()
    run_num = p.add_run(f"  {num}  ")
    run_num.bold = True
    run_num.font.size = Pt(10)
    run_text = p.add_run(f"  {text}")
    run_text.font.size = Pt(11)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = cell_text
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    return table

# ============================================================
# COVER PAGE
# ============================================================
for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("PT KERETA API INDONESIA (PERSERO)")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(100, 100, 100)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("🚆")
run.font.size = Pt(60)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("PANDUAN PENGGUNA")
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(100, 100, 100)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("WPCL")
run.font.size = Pt(40)
run.bold = True
run.font.color.rgb = RGBColor(233, 69, 96)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Workstation Passenger Check List")
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Versi 1.0  •  Tahun 2026  •  Depo Kereta Yogyakarta")
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(150, 150, 150)

doc.add_page_break()

# ============================================================
# DAFTAR ISI
# ============================================================
add_heading_styled("Daftar Isi", level=1)

toc_items = [
    ("A", "Pendahuluan"),
    ("B", "Memulai (Install & Setup)"),
    ("C", "Panduan Petugas Lapangan (Mobile)"),
    ("D", "Panduan Web Admin"),
    ("E", "Cetak Laporan PDF"),
    ("F", "Panduan Manager & Supervisor"),
    ("G", "Referensi & Troubleshooting"),
]

for letter, title in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(f"{letter}.  {title}")
    run.font.size = Pt(12)
    run.bold = True

doc.add_page_break()

# ============================================================
# BAB A: PENDAHULUAN
# ============================================================
add_heading_styled("BAB A — PENDAHULUAN", level=1)

add_heading_styled("Tentang Dokumen Ini", level=2)

add_para("Panduan ini ditujukan untuk seluruh pengguna WPCL: Petugas Lapangan, Supervisor, Manager, dan Administrator. Setiap bagian disusun sesuai peran masing-masing.", size=11)

add_heading_styled("Pengenalan WPCL", level=2)

add_para("WPCL (Workstation Passenger Check List) adalah aplikasi digital untuk mencatat hasil pemeriksaan kondisi sarana kereta api penumpang. Seluruh proses pemeriksaan — dari pengisian checklist hingga cetak laporan PDF — dilakukan secara digital.", size=11)

add_para("Fitur Utama:", bold=True, size=11)
features = [
    "Pemeriksaan Digital — Checklist CCTV, PIDS Luar, PIDS Dalam, WiFi per gerbong tanpa kertas",
    "Locotrack — Pencatatan kondisi lokomotif pendukung secara terstruktur",
    "Auto-Print PDF — Laporan resmi otomatis tercetak setelah penyimpanan data",
    "Mobile-Friendly — Bisa digunakan langsung di lapangan melalui HP Android",
]
for f in features:
    p = doc.add_paragraph(f, style='List Bullet')

add_heading_styled("Alur Singkat", level=2)
add_para("1. Isi Identitas  →  2. Checklist Gerbong  →  3. Locotrack  →  4. Simpan & Cetak PDF", bold=True, size=11)

add_heading_styled("Persyaratan & Akses", level=2)

add_table(
    ["Kebutuhan", "Keterangan"],
    [
        ["HP Android", "Aplikasi mobile WPCL membutuhkan HP Android dengan koneksi internet"],
        ["Koneksi Internet", "WiFi atau jaringan seluler yang stabil"],
        ["Browser (Web Admin)", "Chrome, Firefox, atau Edge untuk akses panel admin dari PC/laptop"],
        ["Printer", "Terhubung ke PC/laptop untuk cetak laporan PDF"],
    ]
)

doc.add_paragraph()

add_heading_styled("Peran Pengguna", level=2)
add_table(
    ["Peran", "Akses", "Fitur Utama"],
    [
        ["Petugas Lapangan", "Mobile App", "Isi checklist, locotrack, cetak PDF"],
        ["Supervisor", "Web Admin", "Review data, edit, cetak ulang"],
        ["Manager", "Web Admin", "Lihat laporan, monitoring"],
        ["Admin", "Web Admin", "Kelola data kereta, sarana, pejabat"],
    ]
)

doc.add_page_break()

# ============================================================
# BAB B: MEMULAI
# ============================================================
add_heading_styled("BAB B — MEMULAI", level=1)

add_heading_styled("Install Aplikasi Mobile", level=2)

add_step(1, "Minta file APK kepada administrator WPCL. File bernama: wpcl_app-release.apk")
add_step(2, "Buka file APK di HP Android. Jika muncul peringatan \"Sumber tidak dikenal\", pilih Izinkan atau Install dari sumber ini.")
add_step(3, "Tunggu proses instalasi selesai (~1 menit). Aplikasi WPCL akan muncul di menu HP.")
add_step(4, "Buka aplikasi WPCL. Halaman \"Identitas Pemeriksaan\" akan terbuka.")

add_para("Tips: Setelah install, pastikan HP terhubung ke internet sebelum membuka aplikasi.", italic=True, size=10, color=(100, 100, 100))

add_heading_styled("Setup Koneksi API", level=2)

add_step(1, "Buka aplikasi WPCL → klik ikon ⚙️ (gear) di pojok kanan atas.")
add_step(2, "Masukkan URL API server: http://<IP_SERVER>:8000/api")
add_step(3, "Contoh URL:")
p = doc.add_paragraph("    Lokal: http://192.168.100.57:8000/api", style='List Bullet')
p = doc.add_paragraph("    Online: https://wpcl-api.vercel.app/api", style='List Bullet')
add_step(4, "Klik \"Simpan\".")

add_para("Jika berhasil, aplikasi akan menampilkan daftar kereta di halaman \"Identitas Pemeriksaan\".", bold=True, size=11)

add_heading_styled("Cara Mencari IP Server", level=2)

add_para("IP server adalah alamat PC/laptop yang menjalankan backend WPCL. HP harus 1 jaringan (WiFi) yang sama.", size=11)

add_step(1, "Buka CMD (Command Prompt) di PC/laptop server.")
add_step(2, "Ketik perintah: ipconfig")
add_step(3, "Cari baris \"IPv4 Address\" — contoh: 192.168.100.57")
add_step(4, "Masukkan IP tersebut ke aplikasi HP: http://192.168.100.57:8000/api")

add_heading_styled("Akses Online via Cloudflare Tunnel", level=2)

add_para("Untuk akses dari luar jaringan (tanpa 1 WiFi yang sama), gunakan Cloudflare Tunnel:", size=11)

add_step(1, "Di server, jalankan: cloudflared tunnel --url http://localhost:8000")
add_step(2, "Copy link yang muncul (contoh: https://xxxx.trycloudflare.com)")
add_step(3, "Masukkan ke aplikasi HP: https://xxxx.trycloudflare.com/api")

add_para("Perhatian: Link Cloudflare Tunnel berubah setiap kali server di-restart. Set ulang URL di aplikasi HP jika server di-restart.", bold=True, size=10, color=(200, 100, 0))

doc.add_page_break()

# ============================================================
# BAB C: PANDUAN PETUGAS LAPANGAN
# ============================================================
add_heading_styled("BAB C — PANDUAN PETUGAS LAPANGAN", level=1)

add_heading_styled("Alur Kerja Pemeriksaan", level=2)
add_para("1. Identitas Pemeriksaan  →  2. Checklist Pemeriksaan  →  3. Pemeriksaan Locotrack  →  4. Simpan & Cetak PDF", bold=True, size=11)

add_heading_styled("Langkah 1: Isi Identitas Pemeriksaan", level=2)

add_step(1, "Buka aplikasi WPCL → halaman \"Identitas Pemeriksaan\" akan terbuka.")
add_step(2, "Isi field berikut:")

add_table(
    ["Field", "Keterangan", "Wajib?"],
    [
        ["Nama Petugas", "Nama lengkap petugas", "Ya"],
        ["NIP", "Nomor Induk Pegawai", "Ya"],
        ["No Kereta", "Pilih dari dropdown", "Ya"],
        ["Tanggal", "Tanggal pemeriksaan", "Ya"],
        ["Business Area", "Kode area bisnis", "Opsional"],
        ["No Referensi", "Nomor referensi dokumen", "Opsional"],
    ]
)

doc.add_paragraph()
add_step(3, "Klik tombol \"Selanjutnya\" → masuk ke halaman Checklist Pemeriksaan.")

add_para("Tips: Pastikan data kereta sudah diinput oleh admin melalui web admin sebelum petugas melakukan pemeriksaan.", italic=True, size=10, color=(100, 100, 100))

add_heading_styled("Langkah 2: Checklist Pemeriksaan", level=2)

add_para("Setiap gerbong diperiksa untuk 4 kategori. Pilih kondisi B/R/T:", size=11)

add_table(
    ["Status", "Keterangan"],
    [
        ["B (Baik)", "Pemeriksaan normal"],
        ["R (Rusak)", "Ditemukan kerusakan"],
        ["T (Tiada)", "Tidak ada/tidak terpasang"],
    ]
)

doc.add_paragraph()
add_para("4 Kategori Pemeriksaan:", bold=True, size=11)
add_table(
    ["Kategori", "Yang Diperiksa"],
    [
        ["CCTV", "Kondisi kamera pengawas di dalam kereta"],
        ["PIDS Luar", "Passenger Information Display System bagian luar kereta"],
        ["PIDS Dalam", "Passenger Information Display System bagian dalam kereta"],
        ["WiFi", "Kondisi jaringan WiFi di dalam kereta"],
    ]
)

doc.add_paragraph()
add_step(1, "Untuk setiap gerbong, pilih kondisi B, R, atau T untuk masing-masing kategori.")
add_step(2, "Isi kolom \"Keterangan\" jika ada catatan tambahan (opsional).")
add_step(3, "Klik \"Selanjutnya\" untuk lanjut ke gerbong berikutnya, atau ke halaman Locotrack jika semua gerbong sudah diperiksa.")

add_para("Catatan: Jumlah gerbong mengikuti jumlah sarana yang terdaftar di data kereta. Tambah gerbong via web admin jika diperlukan.", italic=True, size=10, color=(100, 100, 100))

add_heading_styled("Langkah 3: Pemeriksaan Locotrack", level=2)

add_para("Halaman Locotrack mencatat kondisi lokomotif pendukung.", size=11)

add_step(1, "Pilih Status locotrack: B (Baik) / R (Rusak) / T (Tiada)")
add_step(2, "Isi No Sarana — nomor sarana lokomotif (contoh: KA-001)")
add_step(3, "Isi ID Loco — identitas lokomotif")
add_step(4, "Isi Keterangan — kondisi atau catatan tambahan")

add_para("Field di atas bersifat opsional. Isi sesuai kondisi di lapangan.", italic=True, size=10, color=(100, 100, 100))

add_heading_styled("Catatan Keseluruhan", level=2)

add_step(1, "Scroll ke bawah halaman Locotrack → temukan kolom \"Catatan\".")
add_step(2, "Ketik catatan keseluruhan untuk seluruh pemeriksaan (opsional).")

add_heading_styled("Langkah 4: Simpan & Cetak PDF", level=2)

add_step(1, "Setelah semua data diisi, klik tombol \"Simpan\" di bagian bawah halaman.")
add_step(2, "Proses otomatis:")
p = doc.add_paragraph("    • Data dikirim ke server & tersimpan di database", style='List Bullet')
p = doc.add_paragraph("    • PDF laporan terbuka untuk preview", style='List Bullet')
p = doc.add_paragraph("    • PDF otomatis tercetak (auto-print)", style='List Bullet')
add_step(3, "Aplikasi kembali ke halaman utama \"Identitas Pemeriksaan\".")

add_para("Pemeriksaan berhasil tersimpan! Lihat riwayat pemeriksaan di menu \"Riwayat\" atau di web admin.", bold=True, size=11)

add_heading_styled("Melihat Riwayat Pemeriksaan", level=2)

add_step(1, "Dari halaman utama, tap tombol \"Riwayat\" di bagian bawah layar.")
add_step(2, "Daftar semua pemeriksaan yang sudah dilakukan akan ditampilkan.")
add_step(3, "Tap salah satu untuk melihat detail pemeriksaan.")

add_heading_styled("Pengaturan Server", level=2)

add_step(1, "Tap ikon ⚙️ (gear) di pojok kanan atas halaman utama.")
add_step(2, "Ubah URL API jika server berganti IP atau menggunakan Cloudflare Tunnel.")
add_step(3, "Klik \"Simpan\".")

add_para("Jika URL salah, aplikasi tidak bisa mengambil data dari server. Pastikan IP dan port benar.", bold=True, size=10, color=(200, 100, 0))

doc.add_page_break()

# ============================================================
# BAB D: WEB ADMIN
# ============================================================
add_heading_styled("BAB D — PANDUAN WEB ADMIN", level=1)

add_heading_styled("Akses Web Admin", level=2)

add_step(1, "Buka browser (Chrome / Firefox / Edge) di PC/laptop.")
add_step(2, "Akses URL: http://<IP_SERVER>:8000")
add_para("Contoh: http://localhost:8000 atau http://192.168.100.57:8000", size=11)
add_step(3, "Halaman Dashboard akan terbuka — menampilkan ringkasan jumlah kereta, sarana, dan pemeriksaan.")

add_para("Web admin tidak memerlukan login. Pastikan akses terbatas untuk jaringan internal saja.", italic=True, size=10, color=(100, 100, 100))

add_heading_styled("Menu Web Admin", level=2)

add_table(
    ["Menu", "Fungsi"],
    [
        ["Dashboard", "Ringkasan data (jumlah kereta, sarana, pemeriksaan)"],
        ["Kereta", "Kelola data kereta (tambah, edit, hapus)"],
        ["Sarana", "Kelola gerbong per kereta"],
        ["Pejabat", "Kelola data pejabat penandatangan"],
        ["Pemeriksaan", "Lihat, edit, cetak ulang laporan"],
    ]
)

doc.add_paragraph()

add_heading_styled("Kelola Data Kereta", level=2)

add_step(1, "Klik menu \"Kereta\" di sidebar kiri.")
add_step(2, "Daftar kereta akan ditampilkan dalam tabel.")
add_step(3, "Tambah Kereta: Klik tombol \"Tambah Kereta\" → isi form (No KA, Nama KA, Relasi, Jam) → Simpan.")
add_step(4, "Edit Kereta: Klik ikon ✏️ pada baris kereta → ubah data → Simpan.")
add_step(5, "Hapus Kereta: Klik ikon 🗑️ pada baris kereta → konfirmasi hapus.")

add_heading_styled("Kelola Data Sarana (Gerbong)", level=2)

add_step(1, "Di halaman Kereta, klik nama kereta → halaman detail sarana terbuka.")
add_step(2, "Tambah Sarana: Klik \"Tambah Sarana\" → isi:")
p = doc.add_paragraph("    • Kode Sarana (contoh: G1, G2, G3)", style='List Bullet')
p = doc.add_paragraph("    • Nomor Sarana", style='List Bullet')
p = doc.add_paragraph("    • Seri Sarana (opsional)", style='List Bullet')
p = doc.add_paragraph("    • Depo Induk (opsional)", style='List Bullet')
add_step(3, "Klik Simpan. Sarana akan muncul di dropdown saat petugas mengisi checklist.")

add_para("Wajib: Data sarana harus diinput terlebih dahulu sebelum petugas bisa melakukan pemeriksaan. Jumlah sarana = jumlah gerbong yang akan diperiksa.", bold=True, size=10, color=(200, 100, 0))

add_heading_styled("Kelola Data Pejabat", level=2)

add_step(1, "Klik menu \"Pejabat\" di sidebar.")
add_step(2, "Tambah Pejabat: Klik \"Tambah Pejabat\" → isi: Nama, NIP, Jabatan → Simpan.")
add_step(3, "Pejabat yang aktif akan ditampilkan di bagian tanda tangan PDF laporan.")

add_heading_styled("Mengedit No Dokumen & Versi", level=2)

add_step(1, "Di halaman daftar pemeriksaan, klik ikon 📄 pada baris pemeriksaan.")
add_step(2, "Modal akan muncul → isi No Dokumen dan Versi → Simpan.")

add_para("No Dokumen dan Versi akan ditampilkan di header PDF laporan.", italic=True, size=10, color=(100, 100, 100))

doc.add_page_break()

# ============================================================
# BAB E: CETAK LAPORAN PDF
# ============================================================
add_heading_styled("BAB E — CETAK LAPORAN PDF", level=1)

add_heading_styled("3 Cara Cetak PDF", level=2)

add_table(
    ["Cara", "Langkah"],
    [
        ["Dari Aplikasi Mobile", "PDF otomatis tercetak setelah klik \"Simpan\" di halaman Locotrack"],
        ["Dari Web Admin", "Klik \"Cetak\" pada baris pemeriksaan di halaman data pemeriksaan"],
        ["Direct URL", "Buka: http://server/pemeriksaan/{id}/print"],
    ]
)

doc.add_paragraph()

add_heading_styled("Struktur Laporan PDF", level=2)

add_table(
    ["Bagian", "Isi"],
    [
        ["Header", "No Dokumen, Versi, Judul Laporan"],
        ["Data Petugas", "Nama, NIP, Tanggal Pemeriksaan"],
        ["Data Kereta", "No KA, Nama KA, Relasi, Jam"],
        ["Tabel Pemeriksaan", "No | CCTV | PIDS Luar | PIDS Dalam | WiFi | Keterangan"],
        ["Locotrack", "No Sarana | Locotrack ID | Kondisi | Keterangan"],
        ["Catatan", "Catatan keseluruhan"],
        ["Tanda Tangan", "Nama, NIP, Jabatan Pejabat"],
    ]
)

doc.add_paragraph()
add_para("PDF laporan siap disimpan, dibagikan, atau dicetak sebagai dokumentasi resmi.", bold=True, size=11)

doc.add_page_break()

# ============================================================
# BAB F: PANDUAN MANAGER & SUPERVISOR
# ============================================================
add_heading_styled("BAB F — PANDUAN MANAGER & SUPERVISOR", level=1)

add_heading_styled("Monitoring Pemeriksaan", level=2)

add_step(1, "Buka web admin → menu \"Pemeriksaan\".")
add_step(2, "Filter berdasarkan tanggal, kereta, atau petugas untuk mencari pemeriksaan tertentu.")
add_step(3, "Klik \"Lihat\" untuk melihat detail pemeriksaan lengkap.")
add_step(4, "Klik \"Cetak\" untuk mengunduh laporan PDF.")

add_heading_styled("Mengecek Kondisi Sarana", level=2)

add_step(1, "Di halaman \"Pemeriksaan\", lihat kolom \"Locotrack\" untuk melihat kondisi lokomotif.")
add_step(2, "Klik detail pemeriksaan untuk melihat kondisi per gerbong (CCTV, PIDS, WiFi).")

add_heading_styled("Dashboard Ringkasan", level=2)

add_para("Halaman Dashboard menampilkan:", size=11)
add_table(
    ["Menu", "Keterangan"],
    [
        ["Jumlah Kereta", "Total kereta yang terdaftar di sistem"],
        ["Jumlah Sarana", "Total gerbong yang terdaftar"],
        ["Jumlah Pemeriksaan", "Total pemeriksaan yang sudah dilakukan"],
        ["Jumlah Pejabat", "Total pejabat yang terdaftar"],
    ]
)

doc.add_page_break()

# ============================================================
# BAB G: TROUBLESHOOTING
# ============================================================
add_heading_styled("BAB G — REFERENSI & TROUBLESHOOTING", level=1)

add_heading_styled("Masalah Umum & Solusi", level=2)

add_table(
    ["Masalah", "Penyebab", "Solusi"],
    [
        ["Aplikasi tidak bisa connect", "URL API salah atau server mati", "Cek URL di ⚙️ → pastikan IP & port benar → pastikan server jalan"],
        ["\"Gagal mengambil data kereta\" (401)", "Token API tidak cocok", "Hubungi admin untuk memperbarui token di server"],
        ["\"Gagal mengambil data kereta\" (404)", "Endpoint tidak ditemukan", "Pastikan URL diakhiri /api (bukan /)"],
        ["Pilihan kereta kosong", "Belum ada data kereta", "Input data kereta via web admin"],
        ["\"Pilih Kereta\" tidak bisa diklik", "Tidak ada sarana di kereta", "Input data sarana/gerbong via web admin"],
        ["PDF tidak tercetak", "Printer tidak terhubung", "Cek koneksi printer → pastikan printer online"],
        ["Web admin tidak bisa dibuka", "Server belum jalan", "Jalankan server: php artisan serve"],
        ["Error \"could not find driver\"", "PHP extension belum aktif", "Buka Laragon → PHP → Extensions → centang pdo_mysql"],
        ["Link Cloudflare tidak bisa diakses", "Link expired (server di-restart)", "Restart Cloudflare tunnel → update URL di HP"],
    ]
)

doc.add_paragraph()

add_heading_styled("Kontak Dukungan", level=2)
add_para("Untuk bantuan teknis, hubungi administrator sistem WPCL di Depo Kereta Yogyakarta.")
add_para("Email: admin-wpcl@kai.co.id", bold=True)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Terima Kasih")
run.font.size = Pt(18)
run.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Semoga WPCL membantu pekerjaan sehari-hari")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(100, 100, 100)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("WPCL v1.0 — Depo Kereta Yogyakarta — KAI 2026")
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(150, 150, 150)

# SAVE
output_path = r"C:\Users\Daniel\OneDrive\Desktop\Project\wpcl-api\Panduan_Pengguna_WPCL.docx"
doc.save(output_path)
print(f"File saved: {output_path}")
