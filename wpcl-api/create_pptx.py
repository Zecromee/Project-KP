from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
DARK = RGBColor(10, 22, 40)
ACCENT = RGBColor(233, 69, 96)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(150, 150, 150)
LIGHT_BG = RGBColor(240, 243, 247)
DARK_BLUE = RGBColor(15, 52, 96)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name='Segoe UI'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf

def add_para(tf, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, space_before=6):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = 'Segoe UI'
    p.alignment = align
    p.space_before = Pt(space_before)
    return p

def add_shape_bg(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_circle(slide, left, top, size, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

# ====== SLIDE 1: COVER ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text(slide, 0.5, 0.5, 12, 0.5, 'PT KERETA API INDONESIA (PERSERO)', size=11, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 1.5, 12, 1, '🚆', size=72, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 2.8, 12, 0.6, 'Panduan Pengguna', size=22, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 3.4, 12, 1, 'WPCL', size=64, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 4.4, 12, 0.6, 'Workstation Passenger Check List', size=20, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 5.5, 12, 0.5, 'Versi 1.0  •  Tahun 2026  •  Depo Kereta Yogyakarta', size=12, color=GRAY, align=PP_ALIGN.CENTER)

# ====== SLIDE 2: DAFTAR ISI ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 12, 0.7, 'PANDUAN PENGGUNA WPCL', size=12, color=GRAY, align=PP_ALIGN.LEFT)
add_text(slide, 8, 0.25, 5, 0.7, 'Daftar Isi', size=12, color=GRAY, align=PP_ALIGN.RIGHT)

add_text(slide, 0.8, 1.5, 11, 0.7, 'Daftar Isi', size=32, color=DARK, bold=True)

toc = [('A', 'Pendahuluan'), ('B', 'Memulai (Install & Setup)'), ('C', 'Panduan Petugas Lapangan (Mobile)'), ('D', 'Panduan Web Admin'), ('E', 'Cetak Laporan PDF'), ('F', 'Panduan Manager & Supervisor'), ('G', 'Referensi & Troubleshooting')]
for i, (letter, title) in enumerate(toc):
    y = 2.3 + i * 0.6
    circle = add_circle(slide, 1.0, y, 0.4, ACCENT)
    add_text(slide, 1.05, y + 0.02, 0.4, 0.4, letter, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.7, y + 0.02, 8, 0.4, title, size=16, color=DARK)

add_text(slide, 0.6, 6.8, 12, 0.5, 'Depo Kereta Yogyakarta — KAI', size=11, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 3: BAB A SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'A', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'PENDAHULUAN', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Tentang WPCL, Persyaratan, dan Alur Kerja', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 4: TENTANG WPCL ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB A — PENDAHULUAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Apa itu WPCL?', size=28, color=DARK, bold=True)
add_text(slide, 0.6, 2.2, 12, 0.8, 'WPCL (Workstation Passenger Check List) adalah aplikasi digital untuk mencatat hasil pemeriksaan kondisi sarana kereta api penumpang. Seluruh proses pemeriksaan dilakukan secara digital.', size=16, color=RGBColor(80,80,80))

features = [('🔍', 'Pemeriksaan Digital', 'Checklist CCTV, PIDS Luar, PIDS Dalam, WiFi per gerbong'), ('🔧', 'Locotrack', 'Pencatatan kondisi lokomotif pendukung'), ('📄', 'Auto-Print PDF', 'Laporan resmi otomatis tercetak'), ('📱', 'Mobile-Friendly', 'Bisa digunakan langsung di lapangan via HP')]
for i, (icon, title, desc) in enumerate(features):
    col = i % 2
    row = i // 2
    x = 0.8 + col * 6
    y = 3.3 + row * 1.8
    add_shape_bg(slide, x, y, 5.5, 1.5, LIGHT_BG)
    add_text(slide, x + 0.3, y + 0.2, 1, 0.8, icon, size=32)
    add_text(slide, x + 1.5, y + 0.2, 3.5, 0.5, title, size=16, color=DARK, bold=True)
    add_text(slide, x + 1.5, y + 0.7, 3.8, 0.7, desc, size=13, color=GRAY)

# ====== SLIDE 5: ALUR KERJA ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB A — PENDAHULUAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Alur Kerja Pemeriksaan', size=28, color=DARK, bold=True)

steps = ['1', '2', '3', '4']
labels = ['Isi Identitas', 'Checklist\nGerbong', 'Loco-\ntrack', 'Simpan &\nCetak PDF']
for i, (s, l) in enumerate(zip(steps, labels)):
    x = 1 + i * 3
    shape = add_shape_bg(slide, x, 3, 2, 2, DARK)
    add_text(slide, x, 3.2, 2, 0.8, s, size=36, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x, 4, 2, 0.8, l, size=14, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 3:
        add_text(slide, x + 2.1, 3.6, 0.8, 0.8, '→', size=28, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

add_text(slide, 0.6, 5.5, 12, 1, 'Menu: Identitas Pemeriksaan → Checklist → Locotrack → Simpan', size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ====== SLIDE 6: PERSYARATAN ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB A — PENDAHULUAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Persyaratan & Akses', size=28, color=DARK, bold=True)

reqs = [('📱', 'HP Android', 'Aplikasi mobile membutuhkan HP Android + internet'), ('🌐', 'Koneksi Internet', 'WiFi atau seluler stabil — HP & server 1 jaringan'), ('💻', 'Browser', 'Chrome/Firefox/Edge untuk akses web admin'), ('🖨️', 'Printer', 'Untuk cetak laporan PDF')]
for i, (icon, title, desc) in enumerate(reqs):
    col = i % 2
    row = i // 2
    x = 0.8 + col * 6
    y = 2.3 + row * 1.8
    add_shape_bg(slide, x, y, 5.5, 1.5, LIGHT_BG)
    add_text(slide, x + 0.3, y + 0.2, 1, 0.8, icon, size=32)
    add_text(slide, x + 1.5, y + 0.2, 3.5, 0.5, title, size=16, color=DARK, bold=True)
    add_text(slide, x + 1.5, y + 0.7, 3.8, 0.7, desc, size=13, color=GRAY)

# ====== SLIDE 7: PERAN ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB A — PENDAHULUAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Peran Pengguna', size=28, color=DARK, bold=True)

roles = [('👷', 'Petugas Lapangan', 'Mobile App', 'Isi checklist, locotrack, cetak PDF'), ('🔍', 'Supervisor', 'Web Admin', 'Review data, edit, cetak ulang'), ('📊', 'Manager', 'Web Admin', 'Lihat laporan, monitoring'), ('⚙️', 'Admin', 'Web Admin', 'Kelola data kereta, sarana, pejabat')]
for i, (icon, role, access, feat) in enumerate(roles):
    x = 0.5 + i * 3.1
    add_shape_bg(slide, x, 2.5, 2.8, 3.5, LIGHT_BG)
    add_text(slide, x, 2.7, 2.8, 0.8, icon, size=40, align=PP_ALIGN.CENTER)
    add_text(slide, x, 3.5, 2.8, 0.5, role, size=16, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x, 4.1, 2.8, 0.4, access, size=12, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, x, 4.6, 2.8, 1, feat, size=12, color=GRAY, align=PP_ALIGN.CENTER)

# ====== SLIDE 8: BAB B SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'B', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'MEMULAI', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Install Aplikasi & Setup Koneksi', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 9: INSTALL APK ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB B — MEMULAI', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Install Aplikasi Mobile (HP Android)', size=28, color=DARK, bold=True)

install_steps = [
    ('1', 'Minta file APK', 'Dapatkan file wpcl_app-release.apk dari administrator'),
    ('2', 'Buka file APK', 'Jika muncul peringatan "Sumber tidak dikenal", pilih Izinkan'),
    ('3', 'Tunggu instalasi', 'Proses install ~1 menit, aplikasi muncul di menu HP'),
    ('4', 'Buka WPCL', 'Halaman "Identitas Pemeriksaan" akan terbuka'),
]
for i, (num, title, desc) in enumerate(install_steps):
    y = 2.3 + i * 1.2
    add_circle(slide, 0.8, y, 0.45, ACCENT)
    add_text(slide, 0.85, y + 0.02, 0.45, 0.45, num, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 4, 0.4, title, size=16, color=DARK, bold=True)
    add_text(slide, 1.5, y + 0.45, 10, 0.5, desc, size=13, color=GRAY)

# ====== SLIDE 10: SETUP API ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB B — MEMULAI', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Setup Koneksi API', size=28, color=DARK, bold=True)

api_steps = [
    ('1', 'Buka WPCL', 'Klik ikon ⚙️ (gear) di pojok kanan atas'),
    ('2', 'Masukkan URL', 'http://<IP_SERVER>:8000/api'),
    ('3', 'Contoh URL', 'Lokal: http://192.168.1.5:8000/api  |  Online: https://wpcl-api.vercel.app/api'),
    ('4', 'Simpan', 'Klik tombol "Simpan" — aplikasi akan mengambil data dari server'),
]
for i, (num, title, desc) in enumerate(api_steps):
    y = 2.3 + i * 1.2
    add_circle(slide, 0.8, y, 0.45, ACCENT)
    add_text(slide, 0.85, y + 0.02, 0.45, 0.45, num, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 4, 0.4, title, size=16, color=DARK, bold=True)
    add_text(slide, 1.5, y + 0.45, 10, 0.5, desc, size=13, color=GRAY)

add_shape_bg(slide, 0.6, 6.2, 12, 0.8, RGBColor(255, 243, 205))
add_text(slide, 0.9, 6.3, 11, 0.6, '⚠️ Pastikan HP dan server 1 Wi-Fi yang sama untuk akses lokal.', size=13, color=RGBColor(133, 100, 4))

# ====== SLIDE 11: IP SERVER ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB B — MEMULAI', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Cara Mencari IP Server', size=28, color=DARK, bold=True)

ip_steps = [
    ('1', 'Buka CMD', 'Buka Command Prompt di PC/laptop server'),
    ('2', 'Ketik perintah', 'ipconfig'),
    ('3', 'Cari IPv4', 'Baris "IPv4 Address" — contoh: 192.168.1.5 atau 192.168.100.57'),
    ('4', 'Masukkan ke HP', 'http://192.168.100.57:8000/api'),
]
for i, (num, title, desc) in enumerate(ip_steps):
    y = 2.3 + i * 1.2
    add_circle(slide, 0.8, y, 0.45, ACCENT)
    add_text(slide, 0.85, y + 0.02, 0.45, 0.45, num, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 4, 0.4, title, size=16, color=DARK, bold=True)
    add_text(slide, 1.5, y + 0.45, 10, 0.5, desc, size=13, color=GRAY)

add_text(slide, 0.6, 5.8, 12, 0.6, 'Akses Online via Cloudflare Tunnel', size=20, color=DARK, bold=True)
add_text(slide, 0.6, 6.4, 12, 0.8, 'Jalankan: cloudflared tunnel --url http://localhost:8000 → copy link → masukkan ke HP', size=13, color=GRAY)

# ====== SLIDE 12: BAB C SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'C', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'PANDUAN PETUGAS LAPANGAN', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Checklist, Locotrack, dan Cetak PDF via Mobile', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 13: IDENTITAS ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB C — PETUGAS LAPANGAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Langkah 1: Isi Identitas Pemeriksaan', size=28, color=DARK, bold=True)

fields = [('Nama Petugas', 'Wajib', '✅'), ('NIP', 'Wajib', '✅'), ('No Kereta', 'Pilih dari dropdown', '✅'), ('Tanggal', 'Tanggal pemeriksaan', '✅'), ('Business Area', 'Opsional', '❌'), ('No Referensi', 'Opsional', '❌')]

add_shape_bg(slide, 0.6, 2.3, 7, 4.5, LIGHT_BG)
for i, (field, ket, wajib) in enumerate(fields):
    y = 2.5 + i * 0.65
    add_text(slide, 1.0, y, 2.5, 0.4, field, size=14, color=DARK, bold=True)
    add_text(slide, 3.5, y, 2.5, 0.4, ket, size=13, color=GRAY)
    add_text(slide, 6.2, y, 0.5, 0.4, wajib, size=14, align=PP_ALIGN.CENTER)

add_shape_bg(slide, 8, 2.3, 4.5, 4.5, DARK)
add_text(slide, 8.3, 2.5, 4, 0.5, 'Setelah semua data diisi:', size=16, color=WHITE, bold=True)
add_text(slide, 8.3, 3.2, 4, 1.5, '• Isi field wajib (Nama, NIP, Kereta, Tanggal)\n• Field opsional boleh dikosongkan\n• Klik "Selanjutnya" untuk lanjut', size=13, color=RGBColor(180,180,180))
add_shape_bg(slide, 8.3, 5.2, 4, 1.2, ACCENT)
add_text(slide, 8.3, 5.3, 4, 1, '→ Selanjutnya', size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# ====== SLIDE 14: CHECKLIST ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB C — PETUGAS LAPANGAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Langkah 2: Checklist Pemeriksaan', size=28, color=DARK, bold=True)

# B/R/T badges
brt = [('✅ B — Baik', RGBColor(40,167,69)), ('❌ R — Rusak', RGBColor(220,53,69)), ('🚫 T — Tiada', RGBColor(255,193,7))]
for i, (label, color) in enumerate(brt):
    x = 0.8 + i * 3.5
    shape = add_shape_bg(slide, x, 2.3, 3, 0.8, color)
    add_text(slide, x, 2.35, 3, 0.7, label, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

add_text(slide, 0.6, 3.4, 12, 0.5, '4 Kategori Pemeriksaan:', size=18, color=DARK, bold=True)

cats = [('CCTV', 'Kamera pengawas di dalam kereta'), ('PIDS Luar', 'Passenger Info Display System luar'), ('PIDS Dalam', 'Passenger Info Display System dalam'), ('WiFi', 'Jaringan WiFi di dalam kereta')]
for i, (cat, desc) in enumerate(cats):
    col = i % 2
    row = i // 2
    x = 0.8 + col * 6
    y = 4.1 + row * 1.2
    add_shape_bg(slide, x, y, 5.5, 1, LIGHT_BG)
    add_text(slide, x + 0.3, y + 0.1, 2, 0.4, cat, size=14, color=DARK, bold=True)
    add_text(slide, x + 0.3, y + 0.5, 5, 0.4, desc, size=12, color=GRAY)

add_text(slide, 0.6, 6.7, 12, 0.5, 'Isi B/R/T per kategori per gerbong → Klik "Selanjutnya" → Lanjut ke gerbong berikutnya', size=13, color=GRAY)

# ====== SLIDE 15: LOCOTRACK ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB C — PETUGAS LAPANGAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Langkah 3: Pemeriksaan Locotrack', size=28, color=DARK, bold=True)

loco_steps = [
    ('1', 'Pilih Status', 'B (Baik) / R (Rusak) / T (Tiada)'),
    ('2', 'Isi No Sarana', 'Nomor sarana lokomotif'),
    ('3', 'Isi ID Loco', 'Identitas lokomotif'),
    ('4', 'Isi Keterangan', 'Kondisi atau catatan tambahan'),
]
for i, (num, title, desc) in enumerate(loco_steps):
    y = 2.3 + i * 1
    add_circle(slide, 0.8, y, 0.4, ACCENT)
    add_text(slide, 0.85, y + 0.01, 0.4, 0.4, num, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 3, 0.4, title, size=15, color=DARK, bold=True)
    add_text(slide, 4.5, y, 6, 0.4, desc, size=13, color=GRAY)

add_text(slide, 0.6, 5.8, 12, 0.5, 'Catatan Keseluruhan', size=18, color=DARK, bold=True)
add_text(slide, 0.6, 6.3, 12, 0.5, 'Scroll ke bawah → Ketik catatan → Klik "Simpan" → PDF otomatis tercetak', size=14, color=GRAY)

# ====== SLIDE 16: SIMPAN PDF ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB C — PETUGAS LAPANGAN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Langkah 4: Simpan & Cetak PDF', size=28, color=DARK, bold=True)

save_steps = [
    ('1', 'Klik "Simpan"', 'Data dikirim ke server dan tersimpan di database'),
    ('2', 'PDF Terbuka', 'Preview laporan otomatis muncul di layar'),
    ('3', 'Auto-Print', 'PDF otomatis tercetak ke printer yang terhubung'),
    ('4', 'Selesai', 'Aplikasi kembali ke halaman "Identitas Pemeriksaan"'),
]
for i, (num, title, desc) in enumerate(save_steps):
    y = 2.3 + i * 1.1
    add_circle(slide, 0.8, y, 0.4, ACCENT)
    add_text(slide, 0.85, y + 0.01, 0.4, 0.4, num, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 3.5, 0.4, title, size=15, color=DARK, bold=True)
    add_text(slide, 5, y, 7, 0.4, desc, size=13, color=GRAY)

add_shape_bg(slide, 0.6, 6.2, 12, 0.8, RGBColor(212, 237, 218))
add_text(slide, 0.9, 6.3, 11, 0.6, '✅ Pemeriksaan berhasil tersimpan! Lihat riwayat di menu "Riwayat" atau di web admin.', size=13, color=RGBColor(21, 87, 36))

# ====== SLIDE 17: BAB D SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'D', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'PANDUAN WEB ADMIN', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Kelola Data Kereta, Sarana, Pejabat, dan Pemeriksaan', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 18: WEB ADMIN ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB D — WEB ADMIN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Akses & Menu Web Admin', size=28, color=DARK, bold=True)

add_shape_bg(slide, 0.6, 2.3, 5.5, 2, DARK)
add_text(slide, 0.9, 2.5, 5, 0.4, 'Cara Akses:', size=16, color=WHITE, bold=True)
add_text(slide, 0.9, 3, 5, 1.2, '1. Buka browser (Chrome/Firefox)\n2. Akses: http://<IP_SERVER>:8000\n3. Dashboard akan terbuka', size=13, color=RGBColor(180,180,180))

menus = [('📊', 'Dashboard', 'Ringkasan data'), ('🚂', 'Kereta', 'Kelola data kereta'), ('🚃', 'Sarana', 'Kelola gerbong'), ('👤', 'Pejabat', 'Kelola pejabat'), ('🔍', 'Pemeriksaan', 'Lihat & cetak laporan')]
for i, (icon, menu, desc) in enumerate(menus):
    x = 6.5
    y = 2.3 + i * 0.8
    add_text(slide, x, y, 0.5, 0.5, icon, size=18)
    add_text(slide, x + 0.6, y, 2, 0.4, menu, size=14, color=DARK, bold=True)
    add_text(slide, x + 3, y, 3, 0.4, desc, size=12, color=GRAY)

# ====== SLIDE 19: KELOLA KERETA ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB D — WEB ADMIN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Kelola Data Kereta & Sarana', size=28, color=DARK, bold=True)

kereta_steps = [
    ('1', 'Klik menu "Kereta"', 'Daftar kereta ditampilkan dalam tabel'),
    ('2', 'Tambah Kereta', 'Klik "Tambah Kereta" → isi No KA, Nama, Relasi, Jam → Simpan'),
    ('3', 'Edit/Hapus', 'Klik ikon ✏️ untuk edit, ikon 🗑️ untuk hapus'),
    ('4', 'Kelola Sarana', 'Klik nama kereta → tambah gerbong (kode, nomor, seri, depo)'),
]
for i, (num, title, desc) in enumerate(kereta_steps):
    y = 2.3 + i * 1.1
    add_circle(slide, 0.8, y, 0.4, ACCENT)
    add_text(slide, 0.85, y + 0.01, 0.4, 0.4, num, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 3.5, 0.4, title, size=15, color=DARK, bold=True)
    add_text(slide, 5, y, 7, 0.4, desc, size=13, color=GRAY)

add_shape_bg(slide, 0.6, 6.2, 12, 0.8, RGBColor(255, 243, 205))
add_text(slide, 0.9, 6.3, 11, 0.6, '⚠️ Data sarana WAJIB diinput sebelum petugas bisa melakukan pemeriksaan.', size=13, color=RGBColor(133, 100, 4))

# ====== SLIDE 20: KELOLA PEJABAT ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB D — WEB ADMIN', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Kelola Pejabat & Edit Pemeriksaan', size=28, color=DARK, bold=True)

pejabat_steps = [
    ('1', 'Kelola Pejabat', 'Menu "Pejabat" → Tambah (Nama, NIP, Jabatan) → Simpan'),
    ('2', 'Pejabat di PDF', 'Pejabat aktif ditampilkan di tanda tangan laporan PDF'),
    ('3', 'Edit Pemeriksaan', 'Menu "Pemeriksaan" → Klik "Lihat" → Edit data'),
    ('4', 'No Dokumen', 'Klik ikon 📄 → isi No Dokumen & Versi → Simpan'),
]
for i, (num, title, desc) in enumerate(pejabat_steps):
    y = 2.3 + i * 1.1
    add_circle(slide, 0.8, y, 0.4, ACCENT)
    add_text(slide, 0.85, y + 0.01, 0.4, 0.4, num, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 3.5, 0.4, title, size=15, color=DARK, bold=True)
    add_text(slide, 5, y, 7, 0.4, desc, size=13, color=GRAY)

# ====== SLIDE 21: BAB E SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'E', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'CETAK LAPORAN PDF', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Struktur & Cara Cetak Laporan', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 22: CETAK PDF ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB E — CETAK LAPORAN PDF', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Cara Cetak PDF', size=28, color=DARK, bold=True)

cara = [('📱', 'Dari Mobile', 'PDF otomatis tercetak\nsetelah klik "Simpan"'), ('💻', 'Dari Web Admin', 'Klik "Cetak" pada\nbaris pemeriksaan'), ('🔗', 'Direct URL', 'Buka: /pemeriksaan/{id}/print'), ('🔄', 'Cetak Ulang', 'Detail pemeriksaan →\nklik "Cetak"')]
for i, (icon, title, desc) in enumerate(cara):
    x = 0.5 + i * 3.1
    add_shape_bg(slide, x, 2.5, 2.8, 2.5, LIGHT_BG)
    add_text(slide, x, 2.7, 2.8, 0.8, icon, size=36, align=PP_ALIGN.CENTER)
    add_text(slide, x, 3.5, 2.8, 0.5, title, size=15, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x, 4.1, 2.8, 0.8, desc, size=12, color=GRAY, align=PP_ALIGN.CENTER)

add_text(slide, 0.6, 5.5, 12, 0.5, 'Struktur Laporan PDF:', size=18, color=DARK, bold=True)
pdf_parts = 'Header (No Dokumen, Versi) → Data Petugas → Data Kereta → Tabel Pemeriksaan → Locotrack → Catatan → Tanda Tangan Pejabat'
add_text(slide, 0.6, 6, 12, 0.5, pdf_parts, size=13, color=GRAY)

# ====== SLIDE 23: BAB F SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'F', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'PANDUAN MANAGER & SUPERVISOR', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Monitoring & Dashboard', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 24: MANAGER ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB F — MANAGER & SUPERVISOR', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Monitoring Pemeriksaan', size=28, color=DARK, bold=True)

mgr_steps = [
    ('1', 'Buka "Pemeriksaan"', 'Di web admin, klik menu Pemeriksaan'),
    ('2', 'Filter Data', 'Berdasarkan tanggal, kereta, atau petugas'),
    ('3', 'Lihat Detail', 'Klik "Lihat" untuk data lengkap per gerbong'),
    ('4', 'Cetak PDF', 'Klik "Cetak" untuk unduh laporan'),
]
for i, (num, title, desc) in enumerate(mgr_steps):
    y = 2.3 + i * 1
    add_circle(slide, 0.8, y, 0.4, ACCENT)
    add_text(slide, 0.85, y + 0.01, 0.4, 0.4, num, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, y, 4, 0.4, title, size=15, color=DARK, bold=True)
    add_text(slide, 5.5, y, 6, 0.4, desc, size=13, color=GRAY)

add_text(slide, 0.6, 5.8, 12, 0.5, 'Dashboard Menampilkan:', size=18, color=DARK, bold=True)
dash = [('🚂 Kereta', '🚃 Sarana', '🔍 Pemeriksaan', '👤 Pejabat')]
for i, item in enumerate(dash[0]):
    add_text(slide, 0.8 + i * 3, 6.3, 2.5, 0.4, item, size=14, color=DARK, bold=True)

# ====== SLIDE 25: BAB G SECTION ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_text(slide, 1, 1.5, 11, 1.5, 'G', size=120, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 3.2, 11, 0.8, 'REFERENSI & TROUBLESHOOTING', size=42, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_text(slide, 1, 4.2, 8, 0.5, 'Masalah Umum & Solusi', size=16, color=GRAY, align=PP_ALIGN.LEFT)

# ====== SLIDE 26: TROUBLESHOOTING ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, 0, 0, 13.333, 1.2, DARK)
add_text(slide, 0.6, 0.25, 6, 0.7, 'BAB G — TROUBLESHOOTING', size=11, color=GRAY)

add_text(slide, 0.6, 1.5, 12, 0.6, 'Masalah Umum & Solusi', size=28, color=DARK, bold=True)

issues = [
    ('Aplikasi tidak connect', 'URL salah / server mati', 'Cek URL di ⚙️ → pastikan IP & port benar'),
    ('Error 401 (Unauthorized)', 'Token tidak cocok', 'Hubungi admin untuk update token'),
    ('Pilihan kereta kosong', 'Belum ada data', 'Input data kereta via web admin'),
    ('"Pilih Kereta" tidak bisa klik', 'Tidak ada sarana', 'Input data sarana/gerbong via web admin'),
    ('PDF tidak tercetak', 'Printer tidak terhubung', 'Cek koneksi printer'),
    ('Error "could not find driver"', 'PHP extension belum aktif', 'Laragon → PHP → Extensions → centang pdo_mysql'),
]

for i, (masalah, penyebab, solusi) in enumerate(issues):
    y = 2.3 + i * 0.8
    if i % 2 == 0:
        add_shape_bg(slide, 0.6, y - 0.05, 12, 0.75, LIGHT_BG)
    add_text(slide, 0.8, y, 3.5, 0.5, masalah, size=12, color=DARK, bold=True)
    add_text(slide, 4.5, y, 2.5, 0.5, penyebab, size=11, color=GRAY)
    add_text(slide, 7.2, y, 5.5, 0.5, solusi, size=11, color=DARK)

# ====== SLIDE 27: PENUTUP ======
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text(slide, 0.5, 1.5, 12, 1, '🚆', size=72, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 2.8, 12, 0.8, 'Terima Kasih', size=42, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 3.8, 12, 0.5, 'Semoga WPCL membantu pekerjaan sehari-hari', size=18, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 5, 12, 0.5, 'WPCL v1.0 — Depo Kereta Yogyakarta — KAI 2026', size=13, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0.5, 5.8, 12, 0.5, 'Dokumen ini untuk seluruh pengguna WPCL', size=12, color=RGBColor(80,80,80), align=PP_ALIGN.CENTER)

# Save
output = r'C:\Users\Daniel\OneDrive\Desktop\Project\wpcl-api\Panduan_Pengguna_WPCL_v2.pptx'
prs.save(output)
print(f'Saved: {output}')
print(f'Total slides: {len(prs.slides)}')
